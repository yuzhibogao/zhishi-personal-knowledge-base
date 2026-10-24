"""Original bytes live in SQLite so database backups include the source documents."""
import hashlib
import io
from pathlib import PurePosixPath
from urllib.parse import quote

from fastapi import HTTPException
from fastapi.responses import Response

TYPES = {
    ".pdf": "application/pdf", ".txt": "text/plain; charset=utf-8",
    ".md": "text/plain; charset=utf-8",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def save(conn, doc_id, filename, raw):
    name = PurePosixPath(filename.replace("\\", "/")).name
    ext = PurePosixPath(name).suffix.lower()
    if ext not in TYPES:
        raise ValueError("不支持的原文件格式")
    conn.execute(
        "INSERT OR REPLACE INTO originals(doc_id,filename,mime,sha256,data) VALUES(?,?,?,?,?)",
        (doc_id, name, TYPES[ext], hashlib.sha256(raw).hexdigest(), raw))


def response(conn, doc_id, download=False):
    row = conn.execute("SELECT * FROM originals WHERE doc_id=?", (doc_id,)).fetchone()
    if not row:
        raise HTTPException(404, "此文档没有保留原文件，请重新关联原文件")
    disposition = "attachment" if download else "inline"
    return Response(bytes(row["data"]), media_type=row["mime"], headers={
        "Content-Disposition": f"{disposition}; filename*=UTF-8''{quote(row['filename'])}",
        "X-Content-Type-Options": "nosniff", "Cache-Control": "no-store"})


def sheets(conn, doc_id):
    import openpyxl
    row = conn.execute("SELECT filename,data FROM originals WHERE doc_id=?", (doc_id,)).fetchone()
    if not row or not row["filename"].lower().endswith(".xlsx"):
        raise HTTPException(400, "此文档没有 Excel 原文件")
    wb = openpyxl.load_workbook(io.BytesIO(row["data"]), read_only=True, data_only=False)
    try:
        return [{
            "name": ws.title,
            "rows": [[str(v) if v is not None else "" for v in row]
                     for row in ws.iter_rows(max_row=min(ws.max_row or 1, 200),
                                             max_col=min(ws.max_column or 1, 40), values_only=True)],
            "truncated": (ws.max_row or 0) > 200 or (ws.max_column or 0) > 40,
        } for ws in wb.worksheets[:20]]
    finally:
        wb.close()
