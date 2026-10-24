"""V0.2 文件解析：统一输出 (标题, 纯文本)。

关键决策：
- 直接使用 pdfplumber / python-docx / openpyxl，而不是 markitdown：
  需要自己控制"扫描件判空报错""表格与段落顺序""Excel 按 sheet 分区"，行为可测。
- 解析失败一律抛 ParseError（带明确中文信息），绝不静默存空内容。
"""
import io
import os
import re

from charset_normalizer import from_bytes

SUPPORTED = (".txt", ".md", ".pdf", ".docx", ".xlsx")
CHUNK_THRESHOLD = 50000  # 超过 5 万字才切片（只作用于切片，不作用于准入）


class ParseError(ValueError):
    """解析失败，信息可直接展示给用户。"""


def decode_bytes(raw: bytes) -> str:
    """编码识别顺序：BOM/UTF-8 严格解码 -> GB18030（兼容 GBK）-> charset_normalizer。

    关键决策：charset_normalizer 对很短的 GBK 文本会误判（已在验证中复现），
    而 UTF-8 严格解码几乎不会误吃 GBK 字节，因此先用确定性方法，再用探测兜底。
    """
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        return raw.decode("gb18030")
    except UnicodeDecodeError:
        pass
    best = from_bytes(raw).best()
    if best is None:
        raise ParseError("无法识别文件编码")
    return str(best)


def _parse_txt(raw: bytes) -> str:
    return decode_bytes(raw)


def _parse_pdf(raw: bytes) -> str:
    import pdfplumber
    try:
        pdf = pdfplumber.open(io.BytesIO(raw))
    except Exception as e:  # 加密/损坏
        raise ParseError(f"PDF 无法打开（可能已加密或损坏）：{e}")
    parts = []
    with pdf:
        for i, page in enumerate(pdf.pages, start=1):
            # Some PDFs paint each glyph twice (e.g. simulated bold). Remove only
            # coincident glyphs; avoid guessing columns or rewriting formulas.
            t = (page.dedupe_chars(tolerance=1).extract_text(x_tolerance=1.5, y_tolerance=3) or "").strip()
            if t:
                parts.append(t)
    text = "\n\n".join(parts).strip()
    if not text:
        raise ParseError("无法提取文字：该 PDF 可能是扫描件，建议 OCR 或手输")
    return text


def _parse_docx(raw: bytes) -> str:
    import docx
    try:
        d = docx.Document(io.BytesIO(raw))
    except Exception as e:
        raise ParseError(f"DOCX 无法打开：{e}")
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for tb in d.tables:
        for row in tb.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ParseError("DOCX 中没有可提取的文字")
    return text


def _parse_xlsx(raw: bytes) -> str:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except Exception as e:
        raise ParseError(f"XLSX 无法打开：{e}")
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append("\t".join(cells).rstrip("\t"))
        if rows:
            parts.append(f"## 工作表：{ws.title}\n" + "\n".join(rows))
    wb.close()
    text = "\n\n".join(parts).strip()
    if not text:
        raise ParseError("XLSX 中没有可提取的内容")
    return text


_DISPATCH = {".txt": _parse_txt, ".md": _parse_txt, ".pdf": _parse_pdf,
             ".docx": _parse_docx, ".xlsx": _parse_xlsx}


def parse_file(filename: str, raw: bytes) -> tuple[str, str]:
    """返回 (标题, 纯文本)。标题默认取文件名（不含扩展名）。"""
    base, ext = os.path.splitext(os.path.basename(filename or "未命名"))
    ext = ext.lower()
    if ext not in _DISPATCH:
        hint = "（旧版 .doc/.xls 不支持，请另存为 .docx/.xlsx）" if ext in (".doc", ".xls") else ""
        raise ParseError(f"不支持的格式 {ext or '无扩展名'}{hint}，支持：{' '.join(SUPPORTED)}")
    text = _DISPATCH[ext](raw).strip()
    if not text:
        raise ParseError("解析结果为空，拒绝入库")
    return base or "未命名", text


def split_chunks(text: str, size: int = 800, overlap: int = 100) -> list[tuple[int, int, str]]:
    """按段落边界切片，相邻切片保留少量重叠。返回 [(char_start, char_end, text)]。

    - 先按空行/换行切成段落，累积到 size 为止；
    - 单个超长段落按句末标点再切，仍超长则硬切；
    - overlap 取上一片末尾字符，避免答案被切在边界。
    """
    text = text.strip()
    if not text:
        return []
    if len(text) <= size:
        return [(0, len(text), text)]
    # 段落及其起点
    units = []
    for m in re.finditer(r"[^\n]+", text):
        seg, base = m.group(), m.start()
        if len(seg) <= size:
            units.append((base, seg))
            continue
        pos = 0
        for s in re.finditer(r"[^。！？!?；;.]+[。！？!?；;.]?", seg):
            piece = s.group()
            while len(piece) > size:  # 硬切兜底
                units.append((base + s.start() + (len(s.group()) - len(piece)), piece[:size]))
                piece = piece[size:]
            if piece:
                units.append((base + s.start() + (len(s.group()) - len(piece)), piece))
    chunks, cur_start, cur = [], None, ""
    for start, seg in units:
        if cur and len(cur) + len(seg) + 1 > size:
            chunks.append((cur_start, cur_start + len(cur), cur))
            tail = cur[-overlap:] if overlap else ""
            cur_start = max(0, cur_start + len(cur) - len(tail))
            cur = tail
        if not cur:
            cur_start = start
        cur = (cur + "\n" + seg) if cur else seg
    if cur.strip():
        chunks.append((cur_start, cur_start + len(cur), cur))
    return chunks
